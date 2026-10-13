"""The Compare workspace (facet.ui.md_compare).

What these tests pin:

* the pure helpers: the descriptors two ModelResults share, the
  parameter-difference statements, the value extraction the basket is
  built from, the differences (scalar, curve, distribution, histogram) and
  the aligned overview rows;
* the basket: a crystal site's row equals the numbers recomputed from its
  SiteResult at the project's v_bond, and a model element's row equals the
  numpy means over that element's atoms of the frame's ValenceTable at the
  model's v_bond;
* the staircase overlay draws one curve per basket row with units on both
  axes and the current v_bond marked; the crystal-vs-glass histogram puts
  the vertical line at the crystal site's value; both render on the
  centre's Figure page;
* two models side by side: the Overview aligns the headline rows and its
  difference column is B − A on a known scalar; a shared Series is overlaid
  with the models' own styles and its Difference toggle equals numpy B − A;
  the Table export round-trips one column per model; the 3D side by side
  holds n_A + n_B atoms with the second block offset by the gap; the tree
  is grouped by question with plain names;
* entering Compare switches the centre, toolbar, Elements dock and column,
  and Back restores all four;
* the View menu action exists on a mnemonic no other View item claims; the
  Sites-list menu, the 3D context menu and the Elements-dock menu all fill
  the basket;
* no verdict word appears in the module or in any text the workspace shows.

Default platform, as the other workspace tests.
"""
from __future__ import annotations

import csv
import dataclasses
import math
import re
import time
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from conftest import dispose  # noqa: E402

HERE = Path(__file__).resolve().parent
DATA = HERE / "data" / "md"
QUARTZ = HERE / "data" / "crystals" / "quartz_SiO2_cod9013321.cif"

VERDICT_WORDS = ("good", "bad", "poor", "excellent", "acceptable", "correct",
                 "incorrect", "wrong", "reliable", "unreliable",
                 "trustworthy", "untrustworthy", "unusable", "should",
                 "proves", "confirms")
VERDICT = re.compile(r"\b(" + "|".join(VERDICT_WORDS) + r")\b",
                     re.IGNORECASE)


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


def _wait(qapp, condition, timeout_s: float = 30.0) -> bool:
    end = time.perf_counter() + timeout_s
    while time.perf_counter() < end:
        qapp.processEvents()
        if condition():
            return True
        time.sleep(0.005)
    qapp.processEvents()
    return bool(condition())


@pytest.fixture(scope="module")
def glass_results():
    """Two runs of the glass analysis on glass.extxyz: one frame, then
    both frames, so their provenances differ in the frames analysed."""
    from facet.core import md_analysis as ma, md_readers

    trajectory = md_readers.read_trajectory(str(DATA / "glass.extxyz"))
    one = ma.analyse(trajectory, ma.AnalysisRequest(
        analyses=("glass",), formers=frozenset({"Si"}), frames=slice(0, 1)))
    both = ma.analyse(trajectory, ma.AnalysisRequest(
        analyses=("glass",), formers=frozenset({"Si"}), frames=slice(0, 2)))
    assert not one.failed and not both.failed
    return one, both


@pytest.fixture(scope="module")
def two_v_bond_results():
    """The glass analysis on glass.extxyz twice, at two bond thresholds."""
    from facet.core import md_analysis as ma, md_readers

    trajectory = md_readers.read_trajectory(str(DATA / "glass.extxyz"))
    low = ma.analyse(trajectory, ma.AnalysisRequest(
        analyses=("glass",), formers=frozenset({"Si"}), v_bond_vu=0.1))
    high = ma.analyse(trajectory, ma.AnalysisRequest(
        analyses=("glass",), formers=frozenset({"Si"}), v_bond_vu=0.4))
    assert not low.failed and not high.failed
    return low, high


@pytest.fixture
def warnings(monkeypatch):
    from PySide6.QtWidgets import QMessageBox

    seen: list[tuple[str, str]] = []

    def record(_parent, title, text, *args, **kwargs):
        seen.append((title, text))
        return QMessageBox.Ok

    monkeypatch.setattr(QMessageBox, "warning", staticmethod(record))
    monkeypatch.setattr(QMessageBox, "information", staticmethod(record))
    return seen


@pytest.fixture
def window(qapp, warnings):
    from facet.ui.preview import PreviewWindow

    w = PreviewWindow()
    yield w
    dispose(w)


def _open_model(qapp, window, path):
    controllers = window.open_md_models([str(path)])
    assert len(controllers) == 1
    controller = controllers[0]
    assert _wait(qapp, lambda: controller.trajectory is not None
                 or controller.problem is not None), "the file was not read"
    assert controller.trajectory is not None, controller.problem
    index = window.project.index_of(controller.entry)
    window.project.set_active(index)
    window.structure_panel.refresh()
    window._show_entry()
    assert _wait(qapp, lambda: controller.entry.frame_view is not None
                 and controller.entry.frame_view.table is not None), \
        "no frame was drawn with its table"
    return controller


def _two_models(qapp, window, results):
    first = _open_model(qapp, window, DATA / "glass.extxyz")
    second = _open_model(qapp, window, DATA / "glass.extxyz")
    first.entry.result, second.entry.result = results
    window._show_compare()
    ws = window._compare_ws
    assert ws is not None and ws.models_list.count() == 2
    return first, second, ws


def _descriptor_of(result, kind):
    for analysis, name, container in result.descriptors():
        if isinstance(container, kind):
            return analysis, name, container
    raise AssertionError(f"no {kind.__name__} in the result")


# ---------------------------------------------------------------------------
# the pure helpers
# ---------------------------------------------------------------------------

def test_common_descriptors_are_those_in_at_least_two_results(glass_results):
    from facet.ui.md_compare import all_descriptors, common_descriptors

    one, both = glass_results
    shared = common_descriptors([one, both])
    everything = {(a, n) for a, n, _c in one.descriptors()} \
        & {(a, n) for a, n, _c in both.descriptors()}
    assert set(shared) == everything and shared

    lonely = SimpleNamespace(descriptors=lambda: iter(
        [("glass", "only here", None)]))
    assert ("glass", "only here") not in common_descriptors([one, lonely])
    assert ("glass", "only here") in all_descriptors([one, lonely])


def test_parameter_differences_state_frames_formers_and_v_bond(
        glass_results):
    from facet.ui.md_compare import parameter_differences

    one, both = glass_results
    assert parameter_differences([("a", one), ("b", one)]) == []
    stated = parameter_differences([("a", one), ("b", both)])
    assert any("frames" in line and "a" in line and "b" in line
               for line in stated)
    moved = dataclasses.replace(
        one, provenance=dataclasses.replace(one.provenance, v_bond_vu=0.31))
    lines = parameter_differences([("a", one), ("b", moved)])
    assert any("v_bond" in line and "0.31" in line for line in lines)
    renamed = dataclasses.replace(
        one, provenance=dataclasses.replace(one.provenance,
                                            formers=frozenset({"P"})))
    lines = parameter_differences([("a", one), ("b", renamed)])
    assert any("formers" in line and "P" in line for line in lines)


def test_crystal_values_equal_the_site_result_they_come_from():
    from facet.core import bv, readers
    from facet.core.coordination import analyse_structure
    from facet.ui.md_compare import crystal_values

    structure = readers.read(QUARTZ)
    result = analyse_structure(structure)[0]
    v = bv.V_BOND_DEFAULT
    row = crystal_values(result, v)
    assert row["cn"] == result.cn_at(v) == 4
    bonded = [c for c in result.contacts
              if c.has_valence and c.valence > v]
    assert row["bvs_vu"] == pytest.approx(
        bv.bvs([c.valence * c.occupancy for c in bonded]))
    assert row["d_mean_ang"] == pytest.approx(
        np.mean([c.distance for c in bonded]))
    plateau = result.current_plateau
    assert row["plateau_decades"] == pytest.approx(plateau.width_decades)
    # moved onto another plateau, the row follows the threshold it was asked
    high = crystal_values(result, 1.2)
    assert high["cn"] == result.cn_at(1.2) < 4


def test_scalar_differences_are_b_minus_a_with_the_relative_change():
    from facet.ui.md_compare import difference, difference_text

    d, rel = difference(2.0, 2.5)
    assert d == pytest.approx(0.5) and rel == pytest.approx(25.0)
    assert difference_text(d, rel, "Å") == "+0.5 Å (+25.0 %)"
    # a percentage row: points, no relative change
    d, rel = difference(0.30, 0.25, percent=True)
    assert d == pytest.approx(-0.05) and math.isnan(rel)
    assert difference_text(d, rel, "1", True) == "-5.0 points"
    # a side lacking a value, or A at 0
    assert all(math.isnan(x) for x in difference(float("nan"), 1.0))
    d, rel = difference(0.0, 1.0)
    assert d == 1.0 and math.isnan(rel)
    assert difference_text(float("nan"), float("nan")) == "—"


def test_series_and_distribution_differences_equal_numpy(glass_results):
    from facet.core.md_stats import Distribution, Series
    from facet.ui.md_compare import (
        difference_figure,
        distribution_difference,
        series_difference,
    )

    one, both = glass_results
    analysis, name, a = _descriptor_of(one, Series)
    b = both.outputs[analysis].tables[name]
    axis, diff = series_difference(a, b)
    assert np.array_equal(axis, a.axis)
    assert np.array_equal(diff, np.asarray(b.mean) - np.asarray(a.mean))
    figure = difference_figure([("A", a), ("B", b)], title=name)
    assert len(figure.series) == 1 and figure.series[0].label == "B − A"
    assert np.array_equal(figure.series[0].mean, diff)
    assert figure.x_label.endswith(")") and figure.y_label.endswith(")")
    # another axis: refused in words, nothing drawn
    other = dataclasses.replace(a, axis=np.asarray(a.axis) * 1.5)
    with pytest.raises(ValueError, match="axes differ"):
        series_difference(a, other)

    analysis, name, da = _descriptor_of(one, Distribution)
    db = both.outputs[analysis].tables[name]
    keys, diff = distribution_difference(da, db)
    assert set(keys) == set(da.keys) | set(db.keys)
    expected = {k: dict(zip(db.keys, db.mean)).get(k, 0.0)
                - dict(zip(da.keys, da.mean)).get(k, 0.0) for k in keys}
    assert diff == pytest.approx([expected[k] for k in keys])
    figure = difference_figure([("A", da), ("B", db)], title=name)
    assert figure.series[0].kind == "bars" and figure.categories is not None
    assert "counts 0" in " ".join(figure.annotations)


def test_overview_rows_align_two_runs_and_state_the_difference(
        two_v_bond_results):
    from facet.ui.md_names import group_titles
    from facet.ui.md_compare import difference_text, overview_rows

    low, high = two_v_bond_results
    rows = overview_rows([("A", low), ("B", high)])
    assert rows and all(len(r["values"]) == 2 for r in rows)
    titles = group_titles()
    order = [titles.index(r["group"]) for r in rows]
    assert order == sorted(order), "rows are in question-group order"

    # a known scalar: the density is the same in both runs
    density = next(r for r in rows if r["key"] == "density")
    assert density["present"] == [True, True]
    assert density["differences"][0][0] == pytest.approx(0.0)
    assert density["difference_texts"][0].startswith("+0")
    # a scalar that moves with v_bond: B − A, as numpy says
    cn = next(r for r in rows if r["key"] == "CN Si (BV)")
    a = low.outputs["glass"].tables["mean CN Si (BV)"].mean
    b = high.outputs["glass"].tables["mean CN Si (BV)"].mean
    assert cn["values"] == pytest.approx([a, b])
    assert cn["differences"][0][0] == pytest.approx(b - a)
    assert cn["difference_texts"][0] == difference_text(
        b - a, 100 * (b - a) / abs(a), cn["unit"])
    assert "bond-valence cut" in cn["label"]
    # the distance cut is its own row wherever a run measured it (a 5-atom
    # frame has no g(r) minimum to cut at, so this fixture may lack it)
    if "mean CN Si (distance)" in low.outputs["glass"].tables:
        assert any(r["key"] == "CN Si (distance)"
                   and "distance cut" in r["label"] for r in rows)

    # a row one side lacks shows nothing there and no difference
    partial = SimpleNamespace(outputs={
        "glass": SimpleNamespace(error=None, tables={
            "density": low.outputs["glass"].tables["density"]})})
    rows = overview_rows([("A", low), ("B", partial)])
    only_a = next(r for r in rows if r["key"] == "CN Si (BV)")
    assert only_a["present"] == [True, False]
    assert only_a["texts"][1] == ""
    assert math.isnan(only_a["differences"][0][0])
    assert only_a["difference_texts"][0] == "—"


def test_comparison_rows_hold_one_column_per_model(two_v_bond_results):
    from facet.ui.md_compare import comparison_rows

    low, high = two_v_bond_results
    rows = comparison_rows([("A", low), ("B", high)])
    assert rows
    for row in rows:
        assert {"A", "B", "B − A", "A mean", "B mean", "unit"} <= set(row)
    density = next(r for r in rows if r["id"] == "density")
    assert density["A mean"] == pytest.approx(
        low.outputs["glass"].tables["density"].mean)
    assert density["descriptor"] == "Mass density"
    # the scope narrows, the filter narrows
    scoped = comparison_rows([("A", low), ("B", high)],
                             scope=("item", "glass", "density"))
    assert [r["id"] for r in scoped] == ["density"]
    assert all("density" in r["descriptor"].lower() or "density" in
               r["id"].lower() for r in comparison_rows(
                   [("A", low), ("B", high)], filter_text="density"))


def test_side_by_side_scene_offsets_the_second_block_by_the_gap():
    from facet.core import md_readers
    from facet.ui.md_compare import side_by_side_scene
    from facet.ui.md_scene import build_md_scene

    trajectory = md_readers.read_trajectory(str(DATA / "glass.extxyz"))
    a = build_md_scene(trajectory.frame(0), show_cell=False)
    b = build_md_scene(trajectory.frame(1), show_cell=False)
    gap = 7.5
    merged, offsets = side_by_side_scene([a, None, b], gap)
    assert merged.n_atoms == a.n_atoms + b.n_atoms
    assert len(offsets) == 3 and np.all(offsets[1] == 0)
    first = merged.atom_position[:a.n_atoms]
    second = merged.atom_position[a.n_atoms:]
    # each block is its scene shifted rigidly along x only
    assert np.allclose(first - a.atom_position, offsets[0], atol=1e-5)
    assert np.allclose(second - b.atom_position, offsets[2], atol=1e-5)
    assert offsets[0][1] == offsets[0][2] == 0.0
    right_of_a = float((first[:, 0] + merged.atom_radius[:a.n_atoms]).max())
    left_of_b = float((second[:, 0]
                       - merged.atom_radius[a.n_atoms:]).min())
    assert left_of_b - right_of_a == pytest.approx(gap, abs=1e-4)
    assert list(np.bincount(merged.structure_of)) == [a.n_atoms, b.n_atoms]
    # the row is centred on x = 0
    lo = float((merged.atom_position[:, 0] - merged.atom_radius).min())
    hi = float((merged.atom_position[:, 0] + merged.atom_radius).max())
    assert lo == pytest.approx(-hi, abs=1e-4)


# ---------------------------------------------------------------------------
# the basket in the window
# ---------------------------------------------------------------------------

def test_the_basket_rows_equal_the_numbers_they_come_from(qapp, window):
    from facet.core import bulk
    from facet.ui.md_compare import model_values

    controller = _open_model(qapp, window, DATA / "glass.extxyz")
    window.load(str(QUARTZ))
    assert window.project.current.kind == "crystal"
    site = window.results[0]
    window._add_site_to_compare(site.site_index)
    ws = window._compare_ws
    assert ws is not None and window.compare_on
    assert window.column_stack.currentWidget() is ws.column
    assert window.centre_stack.currentWidget() is ws.canvas
    window._add_model_to_compare(controller.entry, "O")
    rows = ws.basket_rows()
    assert [r["kind"] for r in rows] == ["crystal", "model"]

    crystal = rows[0]
    assert crystal["site"] == site.label
    assert crystal["cn"] == site.cn_at(window.project.v_bond)

    model = rows[1]
    fv = controller.entry.frame_view
    v_bond = float(controller.strip.v_bond)
    results = bulk.at_threshold(fv.table, v_bond)
    sel = np.asarray(fv.table.elements) == "O"
    assert model["atoms"] == int(sel.sum())
    assert model["cn"] == pytest.approx(
        float(np.asarray(results.cn, float)[sel].mean()))
    bvs = np.asarray(results.bvs_vu, float)[sel]
    assert model["bvs_vu"] == pytest.approx(
        float(bvs[np.isfinite(bvs)].mean()))
    assert model["frame"] == fv.k
    helper = model_values(fv.table, "O", v_bond)
    for key, value in helper.items():
        if key == "element":
            continue
        if isinstance(value, float) and math.isnan(value):
            assert math.isnan(model[key]), key
        else:
            assert model[key] == pytest.approx(value), key
    # the column's table lists one line per reference plus nothing else
    assert ws.basket_table.rowCount() == 2

    # compare stays on across a row switch; Back leaves it
    window.structure_panel.list.setCurrentRow(0)
    qapp.processEvents()
    assert window.project.current is controller.entry
    assert window.column_stack.currentWidget() is ws.column
    assert window.centre_stack.currentWidget() is ws.canvas
    ws.back_button.click()
    assert not window.compare_on
    assert window.column_stack.currentWidget() is controller.column
    assert window.centre_stack.currentWidget() is controller.canvas

    # a v_bond move re-reads the crystal side through the existing signal
    window._on_threshold(1e-3)
    row = ws.basket_rows()[0]
    assert row["cn"] == site.cn_at(1e-3)


def test_the_staircase_overlay_and_the_crystal_vs_glass_line(qapp, window):
    controller = _open_model(qapp, window, DATA / "glass.extxyz")
    window.load(str(QUARTZ))
    site = window.results[0]                        # a Si site of quartz
    assert site.element == "Si"
    window._add_site_to_compare(site.site_index)
    window._add_model_to_compare(controller.entry, "Si")
    ws = window._compare_ws

    figure = ws.staircase_figure()
    assert figure is not None and len(figure.series) == 2
    assert figure.x_label == "v_bond (v.u.)"
    assert figure.y_label == "CN (dimensionless)"
    assert figure.log_x
    labels = [s.label for s in figure.series]
    assert any(site.label in text for text in labels)
    assert any("glass.extxyz" in text and "Si" in text for text in labels)
    v = float(window.project.v_bond)
    assert any(kind == "current" and x == pytest.approx(v)
               for x, _t, kind in figure.vlines)
    # the two curves carry distinct styles, so colour never stands alone
    a, b = (s.style for s in figure.series)
    assert (a.colour, a.marker) != (b.colour, b.marker)

    hist = ws.crystal_vs_glass_figure("CN")
    assert hist is not None and len(hist.series) == 1
    assert hist.series[0].kind == "step"
    assert hist.x_label == "CN (dimensionless)"
    assert "fraction" in hist.y_label
    expected = float(site.cn_at(v))
    assert any(x == pytest.approx(expected) and site.label in text
               for x, text, _k in hist.vlines)

    # both render in the centre, at full width, with the basket as rows
    assert ws.select_site_figure("stairs")
    assert ws.page_shown() == "figure"
    assert ws.view.figure is not None and ws.view.figure.log_x
    assert len(ws.view.rows()) == 2
    assert ws.select_site_figure("hist", "CN")
    assert ws.view.figure is not None
    assert ws.view.figure.series[0].kind == "step"

    # an element with no matching crystal site is not binned
    ws.clear()
    window._add_model_to_compare(controller.entry, "Na")
    assert ws.crystal_vs_glass_figure("CN") is None


# ---------------------------------------------------------------------------
# models side by side
# ---------------------------------------------------------------------------

def test_entering_compare_switches_the_centre_and_back_restores_it(
        qapp, window):
    controller = _open_model(qapp, window, DATA / "glass.extxyz")
    own = controller.widgets[:4]
    stacks = (window.column_stack, window.sites_stack, window.toolbar_stack,
              window.centre_stack)
    assert [s.currentWidget() for s in stacks] == list(own)
    window._show_compare()
    ws = window._compare_ws
    assert [s.currentWidget() for s in stacks] == list(ws.widgets)
    assert window.sites_dock.windowTitle().startswith("Elements")
    assert ws.page_shown() == "overview"
    for name in ("figure", "table", "3d", "overview"):
        ws.show_page(name)
        assert ws.page_shown() == name
        assert ws.page_buttons[name].isChecked()
    ws.back_button.click()
    assert not window.compare_on
    assert [s.currentWidget() for s in stacks] == list(own)
    # from a crystal, Back restores the crystal's widgets
    window.load(str(QUARTZ))
    window._show_compare()
    assert window.centre_stack.currentWidget() is ws.canvas
    ws.back_button.click()
    assert [s.currentIndex() for s in stacks] == [0, 0, 0, 0]


def test_two_models_shared_descriptors_show_side_by_side(qapp, window,
                                                         glass_results):
    from PySide6.QtCore import Qt
    from facet.core.md_stats import Distribution, Series

    one, both = glass_results
    _first, _second, ws = _two_models(qapp, window, glass_results)
    assert all(ws.models_list.item(i).checkState() == Qt.Checked
               for i in range(2))
    # the header states the frame difference, in words, without a verdict
    assert not ws.warning.isHidden()
    assert "frames" in ws.warning.text()
    names = {ws.models_list.item(i).text() for i in range(2)}
    styles = ws.ticked_styles()
    assert (styles[0].colour, styles[0].marker) != (styles[1].colour,
                                                    styles[1].marker)

    analysis, name, _c = _descriptor_of(one, Series)
    assert ws.select_descriptor(analysis, name)
    assert ws.page_shown() == "figure"
    figure = ws.view.figure
    assert figure is not None and len(figure.series) == 2
    assert {s.label for s in figure.series} == names
    assert {s.kind for s in figure.series} == {"line"}
    # the models' own styles, the same as their swatches
    assert [s.style for s in figure.series] == styles
    rows = ws.view.rows()
    assert {r["model"] for r in rows} == names

    analysis, name, _c = _descriptor_of(one, Distribution)
    assert ws.select_descriptor(analysis, name)
    figure = ws.view.figure
    assert figure is not None and len(figure.series) == 2
    assert {s.kind for s in figure.series} == {"bars"}
    assert {s.label for s in figure.series} == names
    assert [s.style for s in figure.series] == styles

    # the filter narrows the tree without touching the data
    before = ws.visible_descriptors()
    assert before
    ws.filter_edit.setText("zzz-no-such-descriptor")
    assert ws.visible_descriptors() == []
    ws.filter_edit.setText("")
    assert ws.visible_descriptors() == before


def test_the_difference_toggle_on_a_series_equals_numpy(qapp, window,
                                                        glass_results):
    from facet.core.md_stats import Series

    one, both = glass_results
    _first, _second, ws = _two_models(qapp, window, glass_results)
    labels = [label for label, _r in ws.ticked_results()]
    analysis, name, a = _descriptor_of(one, Series)
    b = both.outputs[analysis].tables[name]
    assert ws.select_descriptor(analysis, name)
    assert ws.difference_check.isEnabled()
    ws.difference_check.setChecked(True)
    figure = ws.view.figure
    assert figure is not None and len(figure.series) == 1
    assert figure.series[0].label == f"{labels[1]} − {labels[0]}"
    assert np.array_equal(figure.series[0].mean,
                          np.asarray(b.mean) - np.asarray(a.mean))
    assert figure.series[0].style == ws.ticked_styles()[1]
    assert labels[0] in ws.difference_note.text()
    ws.difference_check.setChecked(False)
    assert len(ws.view.figure.series) == 2


def test_the_overview_page_aligns_the_models_in_columns(qapp, window,
                                                        two_v_bond_results):
    from PySide6.QtWidgets import QWidget

    _first, _second, ws = _two_models(qapp, window, two_v_bond_results)
    ws.show_page("overview")
    table = ws.overview_table
    labels = [label for label, _r in ws.ticked_results()]
    headers = [table.horizontalHeaderItem(c).text()
               for c in range(table.columnCount())]
    assert headers[:2] == ["Measurement", "unit"]
    assert headers[2:4] == labels
    assert headers[4] == f"{labels[1]} − {labels[0]}"
    assert table.rowCount() > 3
    texts = [table.item(r, 0).text() for r in range(table.rowCount())
             if table.item(r, 0) is not None]
    assert "Structure" in texts and "Mass density" in texts
    assert any("Coordination of Si — bond-valence cut" == t for t in texts)
    # a swatch on each model column, in the model's colour
    assert not table.horizontalHeaderItem(2).icon().isNull()
    # a bar row holds a drawn bar in each model cell
    bars = [r for r in range(table.rowCount())
            if isinstance(table.cellWidget(r, 2), QWidget)]
    if bars:
        assert isinstance(table.cellWidget(bars[0], 3), QWidget)
    assert "v_bond differs" in ws.overview_header.text()
    assert "v_bond differs" in ws.warning.text()


def test_the_table_export_round_trips_one_column_per_model(
        qapp, window, two_v_bond_results, tmp_path):
    low, high = two_v_bond_results
    _first, _second, ws = _two_models(qapp, window, two_v_bond_results)
    labels = [label for label, _r in ws.ticked_results()]
    ws.show_page("table")
    assert ws.table.rowCount() > 0
    assert set(labels) <= set(ws.table.columns())
    written = ws.export_comparison(tmp_path / "comparison.csv")
    text = written.read_text(encoding="utf-8").splitlines()
    header = [line[2:] for line in text if line.startswith("# ")]
    assert any("glass.extxyz" in line and "v_bond" in line
               and "formers" in line for line in header)
    assert sum(1 for line in header if "glass.extxyz" in line) >= 2
    assert any("v_bond differs" in line for line in header)
    body = [line for line in text if not line.startswith("#")]
    rows = list(csv.DictReader(body))
    assert rows
    for label in labels:
        assert label in rows[0] and f"{label} mean" in rows[0]
    density = next(r for r in rows if r["id"] == "density")
    assert float(density[f"{labels[0]} mean"]) == pytest.approx(
        low.outputs["glass"].tables["density"].mean)
    assert float(density[f"{labels[1]} mean"]) == pytest.approx(
        high.outputs["glass"].tables["density"].mean)
    assert density[f"{labels[1]} − {labels[0]}"].startswith("+0")
    try:
        import openpyxl
    except ImportError:
        return
    book_path = ws.export_comparison(tmp_path / "comparison.xlsx")
    book = openpyxl.load_workbook(book_path, read_only=True)
    assert {"provenance", "overview", "table"} <= set(book.sheetnames)
    sheet = book["table"]
    head = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert set(labels) <= set(head)


def test_the_3d_page_lays_both_frames_side_by_side(qapp, window,
                                                   glass_results):
    first, second, ws = _two_models(qapp, window, glass_results)
    n_a = first.entry.frame_view.scene.n_atoms
    n_b = second.entry.frame_view.scene.n_atoms
    ws.show_page("3d")
    scene = ws.view3d.scene
    assert scene is not None and scene.n_atoms == n_a + n_b
    assert list(np.bincount(scene.structure_of)) == [n_a, n_b]
    note = ws.scene_note.text()
    assert "2 blocks" in note and "frame 0" in note and "Å" in note
    # the gap moves the second block, nothing else
    before = scene.atom_position.copy()
    ws.gap_box.setValue(ws.gap_box.value() + 10.0)
    after = ws.view3d.scene.atom_position
    shift = after[n_a:] - before[n_a:]
    assert np.allclose(shift[:, 0], 5.0, atol=1e-4)     # centred: half each
    assert np.allclose(after[:n_a, 0] - before[:n_a, 0], -5.0, atol=1e-4)
    # the Elements dock lists both models, and a click emphasises one element
    labels = [label for label, _r in ws.ticked_results()]
    assert ws.elements.table.columnCount() == 1 + 2 * len(labels)
    elements = ws.elements.elements()
    assert "O" in elements
    radii_before = ws.view3d.scene.atom_radius.copy()
    ws.elements.choose("O")
    radii = ws.view3d.scene.atom_radius
    kept = np.array([e == "O" for e in ws.view3d.scene.atom_element])
    assert np.allclose(radii[kept], radii_before[kept])
    assert np.all(radii[~kept] < radii_before[~kept])
    assert "O emphasised" in ws.scene_note.text()
    ws.elements.choose("")
    assert np.allclose(ws.view3d.scene.atom_radius, radii_before)


def test_the_tree_is_grouped_by_question_with_plain_names(qapp, window,
                                                          glass_results):
    from facet.ui.md_names import group_titles

    _first, _second, ws = _two_models(qapp, window, glass_results)
    tree = ws.tree
    tops = [tree.topLevelItem(i).text(0)
            for i in range(tree.topLevelItemCount())]
    assert "Structure" in tops
    assert all(t in group_titles() or t == "Site to site" for t in tops)
    structure = next(tree.topLevelItem(i)
                     for i in range(tree.topLevelItemCount())
                     if tree.topLevelItem(i).text(0) == "Structure")
    analyses = [structure.child(k).text(0)
                for k in range(structure.childCount())]
    assert "Coordination, speciation, Qn" in analyses
    node = structure.child(analyses.index("Coordination, speciation, Qn"))
    leaves = [node.child(k).text(0) for k in range(node.childCount())]
    assert "Mass density" in leaves and "density" not in leaves
    kinds = {node.child(k).text(1) for k in range(node.childCount())}
    assert kinds <= {"curve", "distr", "number", "table"}
    counts = {node.child(k).text(2) for k in range(node.childCount())}
    assert counts == {"2 of 2"}
    roles = ws.tree_roles()
    assert ("group", "Structure") in roles
    assert ("item", "glass", "density") in roles


# ---------------------------------------------------------------------------
# entry points
# ---------------------------------------------------------------------------

def test_the_view_menu_action_exists_on_its_own_mnemonic(window):
    from collections import Counter

    action = window.compare_action
    # the bar's action wrappers are kept alive for the whole walk: a menu
    # reached through a collected wrapper is torn down with it (PySide)
    bar_actions = window.menuBar().actions()
    view = next(a.menu() for a in bar_actions if a.text() == "&View")
    items = view.actions()
    assert action in items
    assert action.shortcut().toString() == "M"
    letters = Counter()
    for item in items:
        text = item.text()
        if "&" in text and text.index("&") + 1 < len(text):
            letters[text[text.index("&") + 1].lower()] += 1
    assert letters["m"] == 1
    assert max(letters.values()) == 1


def test_the_three_right_click_menus_fill_the_basket(qapp, window):
    controller = _open_model(qapp, window, DATA / "glass.extxyz")
    window.load(str(QUARTZ))
    qapp.processEvents()

    def action_named(menu, text):
        found = [a for a in menu.actions() if a.text().startswith(text)]
        assert found, [a.text() for a in menu.actions()]
        return found[0]

    site_menu = window._build_site_menu(0)
    action_named(site_menu, "Add to comparison").trigger()
    ws = window._compare_ws
    assert [kind for kind, _e, _k in ws.refs()] == ["crystal"]

    context = window._build_context_menu(0)
    action_named(context, "Add to comparison").trigger()
    assert len(ws.refs()) == 1                 # the same site only once

    assert _wait(qapp, lambda: controller.elements.table.rowCount() > 0)
    element_menu = controller.elements._element_menu(0)
    symbol = controller.elements.table.item(0, 0).text()
    action_named(element_menu, f"Add {symbol} to comparison").trigger()
    kinds = [kind for kind, _e, _k in ws.refs()]
    assert kinds == ["crystal", "model"]
    assert ws.refs()[1][2] == symbol

    # removing the model entry drops its basket row
    row = window.project.index_of(controller.entry)
    window._on_remove(row)
    assert [kind for kind, _e, _k in ws.refs()] == ["crystal"]


# ---------------------------------------------------------------------------
# no verdicts
# ---------------------------------------------------------------------------

def test_no_verdict_word_appears_in_the_module(qapp):
    source = (HERE.parent / "facet" / "ui" / "md_compare.py").read_text(
        encoding="utf-8")
    found = [word for word in VERDICT_WORDS
             if re.search(rf"\b{word}\b", source, re.IGNORECASE)]
    assert not found, found


def test_no_text_the_workspace_shows_carries_a_verdict(qapp, window,
                                                       two_v_bond_results):
    from facet.core.md_stats import Series

    low, _high = two_v_bond_results
    _first, _second, ws = _two_models(qapp, window, two_v_bond_results)
    window.load(str(QUARTZ))
    window._add_site_to_compare(window.results[0].site_index)
    texts = []
    for name in ("overview", "table", "3d", "figure"):
        ws.show_page(name)
        texts += ws.visible_texts()
    analysis, name, _c = _descriptor_of(low, Series)
    ws.select_descriptor(analysis, name)
    ws.difference_check.setChecked(True)
    texts += ws.visible_texts()
    assert len(texts) > 100
    assert [t for t in texts if VERDICT.search(t)] == []
