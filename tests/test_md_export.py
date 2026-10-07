"""Writing an MD analysis run (md_export): CSV per descriptor, XLSX sheets.

What these tests pin:

* a round trip: every CSV file and every sheet, read back, holds the rows of
  its descriptor, every float bit for bit (the CSV writes the shortest text
  that reads back as the same float, not exporters.write_csv's six
  significant figures), NaN as an empty cell;
* every CSV starts with the provenance header the MD prompt lists (file,
  frames used, type map and its source, oxidation states, bond-valence set,
  v_bond, g(r) minima, FACET version, method parameters), and the workbook
  holds the same in its provenance sheet;
* column names carry units, the per-frame columns included; sheet names
  are unique without regard to case, at most 31 characters and free of the
  characters Excel refuses;
* an analysis that produced nothing is listed in the index with its reason;
* each note appears once in a header, and a descriptor over fewer frames
  than its analysis says which;
* without openpyxl the workbook is refused, never written as rounded CSV.
"""
from __future__ import annotations

import csv
import math
import re
from pathlib import Path

import numpy as np
import openpyxl
import pytest

from facet.core import md_analysis as ma, md_export, md_model, md_stats, \
    readers

QUARTZ = Path(__file__).resolve().parent / "data" / "crystals" / \
    "quartz_SiO2_cod9013321.cif"


@pytest.fixture(scope="module")
def result():
    structure = readers.read(QUARTZ)
    base, _ = md_model.supercell_frame(structure, (3, 3, 3))
    rng = np.random.default_rng(3)
    frames = []
    for k in range(2):
        cart = base.cart_ang + rng.normal(0.0, 0.03, base.cart_ang.shape)
        frames.append(md_model.frame_from_arrays(
            base.elements, cart, box_ang=base.box_ang, timestep=k * 100))
    trajectory = md_model.MemoryTrajectory(frames)
    request = ma.AnalysisRequest(
        analyses=("glass", "scattering", "rings", "voronoi", "msd"),
        formers={"Si"}, timestep_fs=1.0,
        scattering=ma.ScatteringOptions(r_window="none",
                                        radiations=("neutron",),
                                        q_grid_max_inv_ang=10.0),
        network=ma.NetworkOptions(ring_criterion="king", ring_max_size=10),
        # a fit window beyond the two frames' lags: msd produces nothing
        dynamics=ma.DynamicsOptions(fit_t_min_ps=1.0, fit_t_max_ps=50.0))
    return ma.analyse(trajectory, request)


def _read_csv(path: Path):
    with open(path, encoding="utf-8", newline="") as handle:
        lines = handle.read().splitlines()
    header = [line[2:] for line in lines if line.startswith("# ")]
    body = [line for line in lines if not line.startswith("#")]
    return header, list(csv.DictReader(body))


def _csv_text_matches(original, text: str) -> bool:
    if original is None:
        return text == ""
    if isinstance(original, (bool, np.bool_)):
        return text == ("yes" if original else "no")
    if isinstance(original, (int, np.integer)):
        return text == str(int(original))
    if isinstance(original, (float, np.floating)):
        if math.isnan(original):
            return text == ""
        if math.isinf(original):
            return text == ("inf" if original > 0 else "-inf")
        return float(text) == float(original)
    if isinstance(original, (tuple, list)):
        return text == "; ".join(str(v) for v in original)
    return text == str(original)


def test_every_csv_reads_back_as_its_descriptor(result, tmp_path):
    paths = md_export.write_csv_files(result, tmp_path)
    _, index = _read_csv(tmp_path / "index.csv")
    listed = [row for row in index if row["file"]]
    assert len(listed) == len(paths) - 2 == sum(
        len(out.tables) for out in result.outputs.values())
    for row in listed:
        container = result.outputs[row["analysis"]].tables[row["descriptor"]]
        expected = md_export.descriptor_rows(container)
        _, rows = _read_csv(tmp_path / row["file"])
        assert len(rows) == len(expected), row["file"]
        assert list(rows[0]) == list(expected[0]), row["file"]
        for got, want in zip(rows, expected, strict=True):
            for key, value in want.items():
                assert _csv_text_matches(value, got[key]), \
                    (row["file"], key, value, got[key])


def test_the_workbook_reads_back_as_the_descriptors(result, tmp_path):
    path = md_export.write_workbook(result, tmp_path / "run.xlsx")
    book = openpyxl.load_workbook(path, read_only=True)
    assert book.sheetnames[:3] == ["provenance", "index", "notes"]
    index = list(book["index"].iter_rows(values_only=True))
    columns = index[0]
    entries = [dict(zip(columns, values, strict=True)) for values in index[1:]]
    sheets = [e for e in entries if e["sheet"]]
    assert len(sheets) == sum(len(o.tables) for o in result.outputs.values())
    for entry in sheets:
        container = result.outputs[entry["analysis"]].tables[
            entry["descriptor"]]
        expected = md_export.descriptor_rows(container)
        values = list(book[entry["sheet"]].iter_rows(values_only=True))
        assert list(values[0]) == list(expected[0]), entry["sheet"]
        assert len(values) - 1 == len(expected)
        for got, want in zip(values[1:], expected, strict=True):
            for cell, value in zip(got, want.values(), strict=True):
                if isinstance(value, float) and math.isnan(value):
                    assert cell is None
                elif isinstance(value, float) and math.isinf(value):
                    assert cell == ("inf" if value > 0 else "-inf")
                elif isinstance(value, float):
                    # openpyxl writes 16 significant digits ('%.16g'): a
                    # relative rounding of at most half a unit of the 16th
                    # digit, 5e-16, held here to 1e-15
                    assert cell == pytest.approx(value, rel=1e-15, abs=0.0), \
                        (entry["sheet"], cell, value)
                elif isinstance(value, (tuple, list)):
                    assert cell == "; ".join(str(v) for v in value)
                else:
                    assert cell == value, (entry["sheet"], cell, value)
    book.close()


def test_every_csv_carries_the_provenance_header(result, tmp_path):
    md_export.write_csv_files(result, tmp_path)
    header, _ = _read_csv(tmp_path / "glass__CN_Si_(BV).csv")
    text = "\n".join(header)
    for field in ("analysis: glass", "descriptor: CN Si (BV)",
                  "program: FACET", "source file:", "frames used: 2: 0, 1",
                  "type map:", "type map source: in memory",
                  "oxidation states: O2- (common), Si4+ (common)",
                  "bond-valence parameters:", "bond threshold v_bond: 0.075",
                  "network formers: Si", "distance cutoff: Si-O",
                  "method parameter: minimum rule = valley",
                  "g(r) first minimum O-Si:",
                  "bulk pair searches per frame (pass 1, pass 2), shared by "
                  "every per-frame analysis: 1, 1"):
        assert field in text, field
    header, _ = _read_csv(tmp_path / "rings__R_C(n)_king_rings.csv")
    assert any(line.startswith("method parameter: network.ring_criterion = "
                               "king") for line in header)


def test_the_workbook_holds_the_run_and_every_analysis_header(result,
                                                              tmp_path):
    path = md_export.write_workbook(result, tmp_path / "run.xlsx")
    book = openpyxl.load_workbook(path, read_only=True)
    lines = [row[0] for row in book["provenance"].iter_rows(values_only=True)
             if row and row[0]]
    for marker in ("== analysis: glass ==", "== analysis: rings ==",
                   "== analysis: voronoi ==", "frames used: 2: 0, 1",
                   "type map source: in memory"):
        assert any(marker in line for line in lines), marker
    book.close()


def test_column_names_carry_units(result):
    glass = result.outputs["glass"].tables
    density = md_export.descriptor_rows(glass["density"])[0]
    assert "mean (g/cm^3)" in density and "std (g/cm^3)" in density
    qn = md_export.descriptor_rows(glass["Qn Si (BV)"])[0]
    assert "mean (fraction)" in qn
    angles = md_export.descriptor_rows(glass["angle Si-O-Si (BV)"])[0]
    assert "bin low (deg)" in angles
    g = md_export.descriptor_rows(glass["g O-Si"])[0]
    assert "r_ang" in g
    s = md_export.descriptor_rows(result.outputs["scattering"].tables[
        "S(Q) neutron"])[0]
    assert "q_inv_ang" in s


def test_sheet_names_are_unique_short_and_legal():
    taken = {"provenance", "index", "notes"}
    names = [md_export._sheet_name("glass", "angle O-Si-O (distance) " * 3,
                                   taken) for _ in range(4)]
    names += [md_export._sheet_name("glass", "a[b]:c*d?e/f\\g", taken),
              md_export._sheet_name("glass", "NOTES", set()),
              md_export._sheet_name("glass", "Index", taken)]
    folded = [n.casefold() for n in names[:5]]
    assert len(set(folded)) == len(folded)
    for name in names:
        assert len(name) <= md_export.SHEET_NAME_MAX
        assert not set(name) & set("[]:*?/\\")
    assert md_export._sheet_name("x", "notes", {"notes"}).casefold() != "notes"


def test_numbers_are_written_in_full(tmp_path):
    value = 0.1 + 0.2                          # 0.30000000000000004
    scalar = md_stats.Scalar("a value", "Å", [value, 1.0 / 3.0])
    assert md_export._csv_value(value) == "0.30000000000000004"
    assert md_export._csv_value(float("nan")) == ""
    assert md_export._csv_value(float("-inf")) == "-inf"
    rows = md_export.descriptor_rows(scalar, per_frame=True)
    assert rows[0]["frame 0 (Å)"] == value and "mean (Å)" in rows[0]


def test_a_failed_analysis_is_listed_with_its_reason(result, tmp_path):
    assert "msd" in result.failed
    md_export.write_csv_files(result, tmp_path)
    _, index = _read_csv(tmp_path / "index.csv")
    failed = [row for row in index if row["analysis"] == "msd"]
    assert failed and "t_max_ps" in failed[0]["not computed"]


def test_export_picks_the_format_from_the_path(result, tmp_path):
    book = md_export.export(result, tmp_path / "sub" / "r.xlsx")
    assert book[0].exists() and book[0].suffix == ".xlsx"
    files = md_export.export(result, tmp_path / "csvdir")
    assert files[0].name == "index.csv" and files[1].name == "provenance.csv"
    with pytest.raises(ValueError, match="directory"):
        md_export.export(result, tmp_path / "one.csv")


def test_per_frame_columns_hold_every_frame_with_its_unit(result):
    """Each per-frame column is named with the unit of its values, as the
    mean is: 140 of 146 files of a --per-frame export had unit-less
    'frame k' columns before."""
    rows = md_export.descriptor_rows(
        result.outputs["glass"].tables["mean CN Si (BV)"], per_frame=True)
    assert {"frame 0 (1)", "frame 1 (1)"} <= set(rows[0])
    g = md_export.descriptor_rows(result.outputs["glass"].tables["g O-Si"],
                                  per_frame=True)[0]
    assert {"frame 0 (1)", "frame 1 (1)"} <= set(g)
    density = md_export.descriptor_rows(
        result.outputs["glass"].tables["density"], per_frame=True)[0]
    assert "frame 1 (g/cm^3)" in density
    qn = md_export.descriptor_rows(
        result.outputs["glass"].tables["Qn Si (BV)"], per_frame=True)[0]
    assert "frame 0 (fraction)" in qn
    for analysis, _, container in result.descriptors():
        for row in md_export.descriptor_rows(container, per_frame=True)[:1]:
            bare = [k for k in row if re.fullmatch(r"(frame|block) -?\d+", k)]
            assert not bare, (analysis, container.name, bare)


def test_a_workbook_without_openpyxl_is_refused_not_rounded(result, tmp_path,
                                                          monkeypatch):
    """Without openpyxl, exporters.write_xlsx writes one six-digit CSV per
    sheet instead of the workbook; write_workbook now refuses and writes
    nothing (the command line reports exit 5)."""
    import sys

    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(ValueError, match="openpyxl"):
        md_export.write_workbook(result, tmp_path / "r.xlsx")
    assert list(tmp_path.iterdir()) == []


def test_each_note_appears_once_in_a_header(result, tmp_path):
    """The provenance notes, the analysis's and the descriptor's overlapped:
    a median header repeated 16 notes, one 93 of 164 lines. Each note is
    now written once."""
    for analysis, descriptor, container in result.descriptors():
        lines = md_export.header_lines(result, analysis, descriptor,
                                       container)
        texts = [line.split(": ", 1)[1] for line in lines
                 if line.startswith(("note: ", "analysis note: ",
                                     "descriptor note: ", "model note: "))]
        assert len(texts) == len(set(texts)), (analysis, descriptor)


def test_a_descriptor_over_fewer_frames_says_which(result):
    """After a cancel in pass 2 the distance descriptors cover fewer frames
    than their analysis's header lists; the header says which."""
    table = ma.Table("shells", ({"a": 1},), frames=(1,))
    lines = md_export.header_lines(result, "glass", "shells", table)
    assert "descriptor frames used: 1: 1 (of the analysis's frames above)" \
        in lines
    whole = md_export.header_lines(
        result, "glass", "CN Si (BV)",
        result.outputs["glass"].tables["CN Si (BV)"])
    assert not any(line.startswith("descriptor frames used") for line in whole)
